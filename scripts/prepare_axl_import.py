"""Converte o layout Controle_Vendas_AXL em importação transacional.
Requer openpyxl. Não modifica o Excel original. Saída contém dados privados.
Uso: python scripts/prepare_axl_import.py arquivo.xlsx pasta-de-saida email-admin
"""
import base64
import collections
import datetime
import hashlib
import json
import pathlib
import re
import sys
import warnings
from decimal import Decimal, ROUND_HALF_UP
import openpyxl
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

def cents(value):
    return int((Decimal(str(value))*100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))

def dates(label):
    label = str(label)
    if re.fullmatch(r'\d{2}/\d{2}/\d{4}', label):
        d = datetime.datetime.strptime(label, '%d/%m/%Y').date().isoformat()
        return d, d, d
    m = re.fullmatch(r'(\d{2})[–-](\d{2})/(\d{2})/(\d{4})', label)
    if m:
        a, b, month, year = map(int, m.groups())
        start, end = datetime.date(year, month, a), datetime.date(year, month, b)
        if start>end: raise ValueError('Período invertido')
        return '', start.isoformat(), end.isoformat()
    if label == 'Sem data precisa': return '', '', ''
    raise ValueError('Data não reconhecida: '+label)

def prepare(path):
    original = pathlib.Path(path).read_bytes()
    w = openpyxl.load_workbook(path, data_only=True)
    rows = []
    occurrences = collections.Counter()
    for i, r in enumerate(w['Vendas AXL'].iter_rows(values_only=True), 1):
        if i==1 or not r[1] or not isinstance(r[3], (int,float)) or not isinstance(r[4], (int,float)): continue
        if int(r[3])!=r[3] or r[3]<=0: raise ValueError('Quantidade inválida na linha '+str(i))
        date, start, end = dates(r[0])
        row = dict(company=str(r[1]).strip(), segment=str(r[2]), plates=int(r[3]), total=cents(r[4]), average_original=str(r[5]), solution=str(r[6]), notes=str(r[7] or ''), date=date, date_start=start, date_end=end, date_label=str(r[0]))
        fingerprint = json.dumps(row, ensure_ascii=False, sort_keys=True)
        occurrences[fingerprint]+=1
        row['key'] = hashlib.sha256((fingerprint+'#'+str(occurrences[fingerprint])).encode()).hexdigest()
        row['sheet_row']=i
        rows.append(row)
    if not rows: raise ValueError('Nenhuma venda encontrada')
    total, plates = sum(r['total'] for r in rows), sum(r['plates'] for r in rows)
    summary = {r[0]:r[1] for r in w['Resumo'].iter_rows(values_only=True) if r[0]}
    assert cents(summary['Faturamento total'])==total, 'Faturamento divergente'
    assert summary['Placas registradas']==plates, 'Quantidade de placas divergente'
    assert summary['Quantidade de vendas']==len(rows), 'Quantidade de vendas divergente'
    company_names = {str(r[1]).strip() for r in w['Empresas'].iter_rows(min_row=6,values_only=True) if r[1] and str(r[0]).isdigit()}
    assert company_names=={r['company'] for r in rows}, 'Cadastro de empresas divergente'
    for sheet in w:
        if re.fullmatch(r'E\d{2}',sheet.title):
            name=sheet.cell(1,1).value
            details=[r for r in sheet.iter_rows(min_row=8,values_only=True) if len(r)>4 and isinstance(r[3],(int,float)) and isinstance(r[4],(int,float))]
            original_rows=[r for r in rows if r['company']==name]
            assert len(details)==len(original_rows), 'Ficha divergente: '+sheet.title
            assert sum(cents(r[4]) for r in details)==sum(r['total'] for r in original_rows), 'Valor de ficha divergente'
    return dict(filename=pathlib.Path(path).name, sha256=hashlib.sha256(original).hexdigest(), count=len(rows), plates=plates, total=total, sheets=w.sheetnames, file_base64=base64.b64encode(original).decode(), rows=rows)

def write(path, output, admin_email):
    p=prepare(path);p['admin_email']=admin_email;folder=pathlib.Path(output);folder.mkdir(parents=True,exist_ok=True);folder.chmod(0o700)
    (folder/'payload.json').write_text(json.dumps(p,ensure_ascii=False),encoding='utf8')
    migration=(pathlib.Path(__file__).resolve().parent.parent/'database/supabase-import.sql').read_text()
    # Uma só transação inclui migração e dados; falha não deixa importação parcial.
    migration=migration.removeprefix('-- Histórico da planilha: aplicar depois das três migrações iniciais.\n').replace('COMMIT;','')
    payload=json.dumps(p,ensure_ascii=False).replace("'","''")
    email_sql=admin_email.replace("'", "''")
    sql=migration+"\nSELECT set_config('request.jwt.claim.sub',(SELECT id::text FROM public.profiles WHERE email='"+email_sql+"' AND role='ADMIN'),true);\nSELECT public.axl_import_history('"+payload+"'::jsonb);\nCOMMIT;\n"
    target=folder/'IMPORTAR_AXL.sql';target.write_text(sql,encoding='utf8');target.chmod(0o600)
    report={k:v for k,v in p.items() if k not in ['rows','file_base64']}
    report['companies']=len({r['company'] for r in p['rows']})
    report['exact_dates']=sum(bool(r['date']) for r in p['rows'])
    report['date_ranges']=sum(bool(r['date_start']) and not r['date'] for r in p['rows'])
    report['unknown_dates']=sum(not r['date_start'] for r in p['rows'])
    (folder/'CONFERENCIA.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    for file in folder.iterdir():
        if file.is_file(): file.chmod(0o600)
    print(json.dumps(report,ensure_ascii=False));return p
if __name__=='__main__':write(sys.argv[1],sys.argv[2],sys.argv[3])
